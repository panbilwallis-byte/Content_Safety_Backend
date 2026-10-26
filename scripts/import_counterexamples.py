"""大模型安全反例库（课题组成果2，32,595 条）批量导入 datasets / dataset_records / tasks。

用法（在 Content_Safety_Backend 目录下）：
    uv run python scripts/import_counterexamples.py "E:/数据治理原型平台/反例库/最终反例库.xlsx"
    uv run python scripts/import_counterexamples.py <文件> --reset   # 覆盖式重导

约定：
- 支持 .xlsx（首行表头）与 .jsonl（键同表头）两种投递格式；
- 训练/测试隔离：每条 payload 标记 usage="test"，训练类任务按该字段排除测试样本；
- 判重：以 sha256(model_output) 为 content_hash，重复行跳过并计入导入任务的 duplicate_count；
- 幂等：同名同版本数据集已存在且未加 --reset 时直接退出，不产生重复数据。
"""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
import uuid
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from sqlalchemy import delete, select

from app.core.database import SessionLocal
from app.core.time import now_shanghai
from app.models.tables import Dataset, DatasetRecord, Task
from app.services.dataset_version_service import purge_dataset_ledger, refresh_raw_version
from app.services.lineage_service import record_ingest_edge

DATASET_NAME = "大模型安全反例库"
DATASET_VERSION = "v1.0.0"
SOURCE_LABEL = "大模型安全反例库（课题组成果2）"
BATCH = 2000


def read_rows(path: Path):
    if path.suffix.lower() == ".jsonl":
        with path.open(encoding="utf-8") as handle:
            for line in handle:
                line = line.strip()
                if line:
                    yield json.loads(line)
        return
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    header = [str(cell).strip() for cell in next(rows)]
    for values in rows:
        if values is None or all(cell is None for cell in values):
            continue
        yield {key: ("" if value is None else str(value).strip())
               for key, value in zip(header, values)}


def build_payload(row: dict) -> dict | None:
    output = row.get("model_output", "")
    if not output:
        return None
    theme = row.get("theme / subtheme", "") or ""
    payload = dict(row)  # 保留原始 24 列，验收时可按任意字段回溯。
    payload.update({
        "text": output,
        "language": "zh",
        "topic_label": theme.split("/")[0].strip() or "综合",
        "usage": "test",
        "source_library": "gu-yyt-lab/MajorDataGovernanceProject 01_最终反例库",
    })
    return payload


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file", type=Path)
    parser.add_argument("--reset", action="store_true", help="删除同名同版本数据集及样本后重导")
    args = parser.parse_args()
    if not args.file.exists():
        print(f"文件不存在：{args.file}")
        return 1

    with SessionLocal() as db:
        existing = db.scalar(select(Dataset).where(
            Dataset.name == DATASET_NAME, Dataset.version == DATASET_VERSION))
        if existing and not args.reset:
            print(f"数据集已存在（id={existing.id}，{existing.name} {existing.version}），"
                  "如需覆盖重导请加 --reset")
            return 1
        if existing and args.reset:
            purge_dataset_ledger(db, existing.id)
            db.execute(delete(DatasetRecord).where(DatasetRecord.dataset_id == existing.id))
            db.delete(existing)
            db.flush()

        dataset = Dataset(
            name=DATASET_NAME, category="安全反例", source_type="industry",
            version=DATASET_VERSION, status="ready",
            description=("大模型安全反例库（课题组成果2）：模型事实性错误输出及其标准事实、"
                         "错误类型与哈希，用于污染样本识别与数据权重判别准确率评估。"),
            metadata_json={"languages": ["zh"], "modalities": ["text"],
                           "source_name": SOURCE_LABEL, "owner": "课题组协作",
                           "usage": "test"})
        db.add(dataset)
        db.flush()

        task_id = f"TR-{uuid.uuid4().hex[:8]}"
        seen: set[str] = set()
        imported = duplicates = 0
        text_bytes = 0
        buffer: list[DatasetRecord] = []
        for row in read_rows(args.file):
            payload = build_payload(row)
            if payload is None:
                continue
            digest = hashlib.sha256(payload["text"].encode("utf-8")).hexdigest()
            if digest in seen:
                duplicates += 1
                continue
            seen.add(digest)
            text_bytes += len(payload["text"].encode("utf-8"))
            buffer.append(DatasetRecord(dataset_id=dataset.id, task_id=task_id,
                                        content_hash=digest, payload=payload))
            imported += 1
            if len(buffer) >= BATCH:
                db.add_all(buffer)
                db.flush()
                buffer.clear()
        if buffer:
            db.add_all(buffer)
            db.flush()

        storage_gb = round(max(text_bytes / 1_000_000_000, 0.01), 3)
        dataset.metadata_json = {**dataset.metadata_json,
                                 "record_count": imported, "storage_gb": storage_gb}
        db.add(Task(
            task_id=task_id, name="反例库批量导入", capability_code="data_ingest",
            trace_id=uuid.uuid4().hex, status="succeeded", source_name=SOURCE_LABEL,
            dataset_name=DATASET_NAME, storage_gb=storage_gb, progress=100,
            success_count=imported, duplicate_count=duplicates, anomaly_count=0,
            input_data={"connectorType": "file", "sourceType": "industry",
                        "createDataset": True, "files": [str(args.file)],
                        "languages": ["zh"], "modalities": ["text"],
                        "sourceName": SOURCE_LABEL, "owner": "课题组协作"},
            config={}, result={"imported": imported, "duplicates": duplicates,
                               "source": "MajorDataGovernanceProject/01_最终反例库"},
            dataset_version=DATASET_VERSION,
            created_at=now_shanghai(), finished_at=now_shanghai()))
        # 批量导入同样是真实接入：冻结原始清单快照并登记谱系边，否则这批数据在版本台账里是黑户 (#13)。
        db.flush()
        refresh_raw_version(db, dataset, task_id=task_id)
        record_ingest_edge(db, task_id=task_id, dataset=dataset)
        db.commit()
        print(f"导入完成：数据集 id={dataset.id}，样本 {imported} 条，"
              f"重复跳过 {duplicates} 条，任务 {task_id}，存储 {storage_gb} GB")
        return 0


if __name__ == "__main__":
    sys.exit(main())
